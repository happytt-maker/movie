#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成答辩/演示用 PPT（docs/演示材料.pptx）。

    python scripts/build_slides.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DST = ROOT / "docs" / "演示材料.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x41)
ACCENT = RGBColor(0x2E, 0x6F, 0x8E)
MUTED = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CN = "微软雅黑"


def text_box(slide, left, top, width, height, text, size=16, bold=False,
             color=NAVY, align=PP_ALIGN.LEFT, spacing=1.1):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = CN
    return box


def slide_base(prs, title: str, kicker: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.16), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if kicker:
        text_box(slide, 0.62, 0.34, 11.5, 0.35, kicker.upper(), size=11, color=MUTED)
    text_box(slide, 0.6, 0.62 if kicker else 0.45, 11.6, 0.8, title, size=28, bold=True)
    line = slide.shapes.add_shape(1, Inches(0.62), Inches(1.32), Inches(11.4), Emu(12700))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT
    line.line.fill.background()
    return slide


def table(slide, rows, left, top, width, col_widths=None, font=12, header_fill=ACCENT):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                   Inches(width), Inches(0.32 * n_rows))
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(Inches(width) * w / total))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(value)
            run.font.size = Pt(font)
            run.font.name = CN
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else LIGHT)
    return shape


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    text_box(s, 1.0, 2.25, 11.3, 1.0, "基于 MovieLens 32M 的", size=40, bold=True, color=WHITE)
    text_box(s, 1.0, 3.1, 11.3, 1.0, "分布式增量电影推荐系统", size=40, bold=True, color=WHITE)
    text_box(s, 1.0, 4.3, 11.3, 0.5, "NDCG@10 = 0.8894　相对热门基线 +20.8%", size=18, color=RGBColor(0x8F, 0xC7, 0xE8))
    text_box(s, 1.0, 5.6, 11.3, 0.4, "大数据技术课程项目 · 2026", size=13, color=MUTED)

    # 2 结论先行
    s = slide_base(prs, "结论先行：排序目标 + 特征融合把 NDCG@10 提升 20.8%", "KEY RESULT")
    cards = [("0.7363", "热门基线"), ("0.8562", "BPR 热启动"), ("0.8894", "集成排序")]
    for i, (value, label) in enumerate(cards):
        left = 0.7 + i * 4.15
        card = s.shapes.add_shape(1, Inches(left), Inches(1.9), Inches(3.7), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT if i < 2 else ACCENT
        card.line.fill.background()
        text_box(s, left + 0.2, 2.15, 3.3, 0.9, value, size=40, bold=True,
                 color=NAVY if i < 2 else WHITE)
        text_box(s, left + 0.2, 3.15, 3.3, 0.5, label, size=15,
                 color=MUTED if i < 2 else WHITE)
    text_box(s, 0.7, 4.5, 11.9, 2.0,
             "· 同一套 BPR 召回，仅把随机初始化换成旧评分模型的 embedding 热启动，就多出 6.6%\n"
             "· LightGBM LambdaRank 融合 9 个特征，在 BPR 之上再提升 3.9%\n"
             "· 全流程一条命令复现：python scripts/run_all.py --config configs/base.yaml",
             size=15, color=NAVY, spacing=1.5)

    # 3 数据与协议
    s = slide_base(prs, "数据与实验协议：严格时间划分，杜绝泄漏", "DATA & PROTOCOL")
    table(s, [
        ["项目", "设定"],
        ["数据规模", "32,000,204 条评分 / 200,948 用户 / 84,432 电影"],
        ["正反馈", "rating ≥ 4"],
        ["划分", "每用户按时间戳 8 : 1 : 1（train 25.5M / valid 3.19M / test 3.29M）"],
        ["候选集", "测试正样本 + 100 个固定负样本（种子 20260907）"],
        ["负样本池", "训练交互 ≥ 5 的条目，排除用户已交互条目"],
        ["指标", "NDCG@10（主）/ Recall@10 / MRR@10"],
    ], 0.7, 1.7, 11.9, col_widths=[2, 9], font=13)

    # 4 架构
    s = slide_base(prs, "系统架构：五层流水线，模块可单独运行", "ARCHITECTURE")
    layers = [
        ("数据层", "时间划分 · 候选集生成 · 质量报告", "prepare_data.py"),
        ("特征层", "条目/用户统计 · 类型亲和度 · 模型分数", "build_features.py"),
        ("算法层", "热门 · LSH · Spark ALS · 流式 · 图 · BPR · LightGBM", "src/*"),
        ("服务层", "按用户切片打分 · Top-10 输出", "top10_*.csv"),
        ("评价层", "NDCG / Recall / MRR / 延迟", "evaluate.py"),
    ]
    for i, (name, desc, file) in enumerate(layers):
        top = 1.65 + i * 1.05
        box = s.shapes.add_shape(1, Inches(0.7), Inches(top), Inches(11.9), Inches(0.88))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
        box.line.color.rgb = LIGHT
        text_box(s, 0.95, top + 0.12, 1.6, 0.5, name, size=16, bold=True, color=ACCENT)
        text_box(s, 2.6, top + 0.14, 7.4, 0.55, desc, size=13, color=NAVY)
        text_box(s, 10.2, top + 0.16, 2.2, 0.5, file, size=11, color=MUTED, align=PP_ALIGN.RIGHT)

    # 5 六任务
    s = slide_base(prs, "六个任务：每项都有可复现结论", "SIX TASKS")
    left_items = [
        ("任务一 数据检查 + 热门", "NDCG 0.7363；四种热度口径对比"),
        ("任务二 相似度 + LSH", "精确 top-K 30.4s；LSH recall 0.148（负面结论）"),
        ("任务三 Spark ALS", "12.9M 正反馈，564s，NDCG 0.790"),
    ]
    right_items = [
        ("任务四 流式增量", "3.29M 事件，78,939 事件/秒，画像 1.73µs/次"),
        ("任务五 二部图", "PageRank≈热门；社区特征 +4.5%"),
        ("任务六 召回-特征-排序", "BPR 0.8562 → 集成 0.8894"),
    ]
    for col, items in enumerate((left_items, right_items)):
        for i, (title, desc) in enumerate(items):
            left = 0.7 + col * 6.1
            top = 1.7 + i * 1.7
            card = s.shapes.add_shape(1, Inches(left), Inches(top), Inches(5.8), Inches(1.45))
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT
            card.line.fill.background()
            text_box(s, left + 0.25, top + 0.14, 5.3, 0.4, title, size=15, bold=True, color=NAVY)
            text_box(s, left + 0.25, top + 0.62, 5.3, 0.7, desc, size=12, color=MUTED)

    # 6 结果
    s = slide_base(prs, "结果：从热门基线到集成排序的完整对比", "RESULTS")
    table(s, [
        ["模型", "NDCG@10", "Recall@10", "相对热门"],
        ["热门基线", "0.7363", "0.7291", "—"],
        ["Spark ALS（20k 抽样）", "0.7896", "—", "—"],
        ["BPR（随机初始化）", "0.8028", "0.7840", "+9.0%"],
        ["BPR（旧模型热启动）", "0.8562", "0.8191", "+16.3%"],
        ["LightGBM 集成排序", "0.8894", "0.8301", "+20.8%"],
    ], 0.7, 1.65, 11.9, col_widths=[4, 2, 2, 2], font=14)
    text_box(s, 0.7, 4.05, 11.9, 0.9,
             "关键：训练目标与指标对齐（BPR 成对损失）比增大模型容量更有效。",
             size=15, bold=True, color=ACCENT)

    # 7 消融
    s = slide_base(prs, "消融实验：BPR 分数是主力，类型特征稳定加分", "ABLATION")
    table(s, [
        ["配置", "NDCG@10", "相对完整模型"],
        ["完整模型（9 特征）", "0.8894", "—"],
        ["去掉 BPR 分数", "0.8417", "−5.4%"],
        ["去掉类型亲和度", "0.8727", "−1.9%"],
    ], 0.7, 1.7, 11.9, col_widths=[4, 2, 2], font=15)
    text_box(s, 0.7, 3.6, 11.9, 2.0,
             "· BPR 分数贡献最大：去掉后退回热门+统计特征水平\n"
             "· 类型亲和度提供 BPR 无法表达的上下文信息，稳定 +1.9%\n"
             "· 时间趋势、贝叶斯均分等特征贡献较弱，可精简",
             size=14, color=NAVY, spacing=1.5)

    # 8 局限与下一步
    s = slide_base(prs, "局限与下一步", "LIMITATIONS")
    text_box(s, 0.7, 1.7, 5.7, 3.5,
             "已知局限\n"
             "· LSH 阈值未调优，recall 仅 0.148\n"
             "· 图实验在 20k 用户子图上完成\n"
             "· 负样本口径由本项目定义\n"
             "· 流式为单机回放，未接 Kafka",
             size=14, color=NAVY, spacing=1.6)
    text_box(s, 6.9, 1.7, 5.7, 3.5,
             "下一步\n"
             "· LSH 按 Jaccard 分布自适应选阈值\n"
             "· Spark GraphX 上做全量社区发现\n"
             "· Structured Streaming 承载实时画像\n"
             "· 序列特征 + 双塔召回",
             size=14, color=NAVY, spacing=1.6)
    text_box(s, 0.7, 5.6, 11.9, 0.8,
             "代码与结果：python scripts/run_all.py 一键复现，results/ 下保存所有指标与 Top-10 文件",
             size=12, color=MUTED)

    prs.save(DST)
    return DST


if __name__ == "__main__":
    path = build()
    print(f"已生成 {path}")