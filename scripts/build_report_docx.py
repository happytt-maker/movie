#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 docs/项目报告.md 转成可编辑的 Word 文档（docs/项目报告.docx）。

依赖 python-docx：推荐用 Office 自带的 Python 运行本脚本，或 pip install python-docx。
    python scripts/build_report_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "项目报告.md"
DST = ROOT / "docs" / "项目报告.docx"

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_rich_text(paragraph, text: str, size: float = 10.5, base_bold: bool = False) -> None:
    """处理 **粗体** 与 `代码`。"""
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, BODY_FONT, size, True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, CODE_FONT, size - 0.5, base_bold)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            run = paragraph.add_run(token)
            set_font(run, BODY_FONT, size, base_bold)


def add_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    header, body = cells[0], [r for r in cells[1:] if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        set_font(run, HEAD_FONT, 9, True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in body:
        cells_row = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cells_row[i].text = ""
            add_rich_text(cells_row[i].paragraphs[0], text, 9)
    doc.add_paragraph()


def main() -> int:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    in_code = False
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        text = "".join(part.strip() for part in buffer)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.line_spacing = 1.4
        add_rich_text(p, text)
        buffer.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            flush()
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(line)
            set_font(run, CODE_FONT, 9)
            i += 1
            continue
        if not line.strip():
            flush()
            i += 1
            continue
        if line.startswith("| "):
            flush()
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            set_font(run, HEAD_FONT, 20, True)
            i += 1
            continue
        if line.startswith("## "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            run = p.add_run(line[3:].strip())
            set_font(run, HEAD_FONT, 15, True)
            i += 1
            continue
        if line.startswith("### "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            run = p.add_run(line[4:].strip())
            set_font(run, HEAD_FONT, 12, True)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            flush()
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, re.sub(r"^\d+\.\s", "", line))
            i += 1
            continue
        if line.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
            i += 1
            continue
        buffer.append(line)
        i += 1

    flush()

    doc.save(DST)
    print(f"已生成 {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())